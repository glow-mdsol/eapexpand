import os
from typing import Dict, Optional
from eapexpand.models.eap import Attribute, Connector, Document, ModelObject

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
import logging

logging.basicConfig()
logger = logging.getLogger(__name__)

HEADERS = [
    "Class",
    "Attribute",
    "Type",
    "Cardinality",
    "Class Note",
    "NCI C-code",
    "Preferred term",
    "Definition",
    "Codelist",
    "External Codelist",
    "Codelist Extensible",
]


def _parse_cardinality(value: Optional[str]) -> tuple[Optional[int], Optional[str]]:
    if not value:
        return None, None
    text = value.strip()
    if ".." in text:
        lower_raw, upper_raw = text.split("..", 1)
        lower = int(lower_raw) if lower_raw.isdigit() else None
        upper = "*" if upper_raw == "*" else (upper_raw if upper_raw.isdigit() else None)
        return lower, upper
    if text == "*":
        return None, "*"
    if text.isdigit():
        number = int(text)
        return number, str(number)
    return None, None


def _merge_cardinality(existing: Optional[str], incoming: Optional[str]) -> Optional[str]:
    if not existing:
        return incoming
    if not incoming or existing == incoming:
        return existing

    e_low, e_high = _parse_cardinality(existing)
    i_low, i_high = _parse_cardinality(incoming)
    if e_low is None or e_high is None or i_low is None or i_high is None:
        return existing

    lower = min(e_low, i_low)
    if e_high == "*" or i_high == "*":
        upper = "*"
    else:
        upper = str(max(int(e_high), int(i_high)))
    return f"{lower}..{upper}"


def _merge_pipe_type(existing: Optional[str], incoming: Optional[str]) -> str:
    values = []
    for value in [existing, incoming]:
        if not value:
            continue
        for part in str(value).split("|"):
            part = part.strip()
            if part and part not in values:
                values.append(part)
    return "|".join(values)


def generate(
    name: str,
    document: Document,
    ct_content: dict,
    codelists: dict,
    output_dir: Optional[str] = "output",
):
    """
    Generates the Excel Representation of the model
    :param name: The name of the model - guides what the output file is called
    :param document: The loaded document
    :param ct_content: The loaded controlled terms
    :param codelists: The loaded codelists
    :param output_dir: The output directory
    """
    print("Generating USDM Excel file")
    packages = {x.package_id: x for x in document.objects if x.object_type == "Package"}
    # remap the code attribute model to the code list model
    # subset by packages
    _partitions = {}
    for _object in document.objects:
        if _object.package_id in packages:
            _package = packages.get(_object.package_id)
            _partitions.setdefault(_package.name, []).append(_object)
        else:
            print("Orphaned Object: ", _object.name, " -> ", _object.package_id)
    cols = {}
    doc = Workbook()

    def write_cell(worksheet, row, column, value, header=False):
        _value = value if value else ""
        columns = cols.setdefault(worksheet.title, {})
        columns[column] = max([len(_value), columns.get(column, 10)])
        cols[worksheet.title] = columns
        worksheet.cell(row, column).value = _value
        worksheet.cell(row, column).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        if header:
            worksheet.cell(row, column).font = Font(bold=True)

    missing_references = {"objects": [], "attributes": [], "connections": []}
    # Write the package sheet
    wkst = doc.active
    wkst.title = "Packages"
    write_cell(wkst, 1, 1, "Package", True)
    write_cell(wkst, 1, 2, "Parent", True)
    write_cell(wkst, 1, 3, "Note", True)
    row_num = 2
    for _package in packages.values():
        write_cell(wkst, row_num, 1, _package.name)
        write_cell(wkst, row_num, 3, _package.note)
        if _package.parent:
            write_cell(wkst, row_num, 2, _package.parent.name)
        row_num += 1
    for _package_name, pobjects in _partitions.items():
        if len(pobjects) == 1:
            continue
        _sheet_name = _package_name[:30]
        # Limit on the Sheet Name length
        sheet = doc.create_sheet(_sheet_name)
        # write the header
        for idx, column in enumerate(HEADERS):
            write_cell(sheet, 1, idx + 1, column, header=True)
        row_num = 2

        for obj in pobjects:
            obj: ModelObject
            _output = {}
            if obj.object_type == "Class":
                _ct_reference = ct_content.get(obj.name)
                if _ct_reference is None:
                    logger.warning(f"Unable to find {obj.name} in CT")
                    continue  # type: Entity
                # super classes, we need to make sure the attributes from the generalization are included, but marked as inherited
                if obj.generalizations:
                    _generalization = document.get_object(
                        obj.generalizations[0].end_object_id
                    )
                else:
                    _generalization = None
                if _ct_reference is None:
                    if obj.name not in missing_references.get("objects"):
                        missing_references["objects"].append(obj.name)
                # _attributes = sorted(
                #     [
                #         attributes[attr_id]
                #         for attr_id in attributes
                #         if attributes[attr_id].object_id == _object_id
                #     ]
                # )
                # write the entity
                write_cell(sheet, row=row_num, column=1, value=obj.name)
                if _ct_reference:
                    if _ct_reference.definition:
                        write_cell(sheet, row=row_num, column=8, value=_ct_reference.definition)
                    if _ct_reference.nci_c_code:
                        write_cell(sheet, row=row_num, column=6, value=_ct_reference.nci_c_code)
                    if _ct_reference.preferred_term:
                        write_cell(
                            sheet, row=row_num, column=7, value=_ct_reference.preferred_term
                        )
                if obj.note:
                    write_cell(sheet, row=row_num, column=5, value=str(obj.note))
                row_num += 1
                # includes object attributes and connections
                for _attribute in obj.attributes:
                    _attribute: Attribute
                    attrib = _output.setdefault(_attribute.name, {})
                    if not attrib:
                        # add a reference to show that the attribute is inherited from the generalization
                        if (
                            _generalization
                            and _attribute in _generalization.object_attributes
                        ):
                            _name = "* " + _attribute.name
                        else:
                            _name = _attribute.name
                        attrib = dict(
                            attribute_name=_name,
                            attribute_type=_attribute.attribute_type,
                            attribute_cardinality=_attribute.cardinality,
                            attribute_note=_attribute.description,
                        )
                        if _ct_reference:
                            _attr_ref = _ct_reference.get_attribute(_attribute.name)
                            if _attr_ref:
                                attrib["definition"] = _attr_ref.definition
                                attrib["c_code"] = _attr_ref.nci_c_code
                                attrib["pref_term"] = _attr_ref.preferred_term
                                if _attr_ref.has_value_list:
                                    # Check if it's an external reference (contains "point")
                                    if _attr_ref.value_list and "point" in _attr_ref.value_list.lower():
                                        attrib["external_value_list"] = _attr_ref.external_code_list
                                        print(f"  → External codelist: {_attr_ref.external_code_list}")
                                    elif _attr_ref.codelist_code:
                                        attrib["codelist"] = _attr_ref.codelist_code
                                        _cl = codelists.get(_attr_ref.codelist_code)
                                        if _cl is not None:
                                            attrib["codelist_extensible"] = "Yes" if _cl.extensible else "No"
                                        print(f"  → Internal codelist: {_attr_ref.codelist_code}")
                            else:
                                missing = (_ct_reference.entity_name, _attribute.name)
                                if missing not in missing_references.get("attributes"):
                                    missing_references["attributes"].append(missing)
                        else:
                            missing = (obj.name, _attribute.name)
                            if missing not in missing_references.get("attributes"):
                                missing_references["attributes"].append(missing)
                    _output[_attribute.name] = attrib
                for outgoing_connection in obj.outgoing_connections:  # type: Connector
                    print(
                        "Adding connection ", outgoing_connection.name, " to ", obj.name
                    )
                    outgoing_connection: Connector
                    _target_name = document.get_object(
                        outgoing_connection.end_object_id
                    ).name
                    attrib = dict(
                        attribute_name=outgoing_connection.name,
                        attribute_type=_target_name,
                        attribute_cardinality=outgoing_connection.dest_card,
                        attribute_note=None,
                    )
                    _attr_ref = _ct_reference.get_attribute(outgoing_connection.name)

                    if _attr_ref:
                        attrib["definition"] = _attr_ref.definition
                        attrib["c_code"] = _attr_ref.nci_c_code
                        attrib["pref_term"] = _attr_ref.preferred_term
                        if _attr_ref.has_value_list:
                            print(f"  Value list for {outgoing_connection.name}: {_attr_ref.value_list}")
                            # Check if it's an external reference (contains "point")
                            if _attr_ref.value_list and "point" in _attr_ref.value_list.lower():
                                attrib["external_value_list"] = _attr_ref.external_code_list
                                print(f"  → External codelist: {_attr_ref.external_code_list}")
                            elif _attr_ref.codelist_code:
                                attrib["codelist"] = _attr_ref.codelist_code
                                _cl = codelists.get(_attr_ref.codelist_code)
                                if _cl is not None:
                                    attrib["codelist_extensible"] = "Yes" if _cl.extensible else "No"
                                print(f"  → Internal codelist: {_attr_ref.codelist_code}")
                    else:
                        missing = (_ct_reference.entity_name, outgoing_connection.name)
                        if missing not in missing_references.get("connections"):
                            missing_references["connections"].append(missing)
                    existing = _output.get(attrib.get("attribute_name"))
                    if existing:
                        existing["attribute_type"] = _merge_pipe_type(
                            existing.get("attribute_type"), attrib.get("attribute_type")
                        )
                        existing["attribute_cardinality"] = _merge_cardinality(
                            existing.get("attribute_cardinality"),
                            attrib.get("attribute_cardinality"),
                        )
                        if not existing.get("definition") and attrib.get("definition"):
                            existing["definition"] = attrib.get("definition")
                        if not existing.get("c_code") and attrib.get("c_code"):
                            existing["c_code"] = attrib.get("c_code")
                        if not existing.get("pref_term") and attrib.get("pref_term"):
                            existing["pref_term"] = attrib.get("pref_term")
                        if not existing.get("codelist") and attrib.get("codelist"):
                            existing["codelist"] = attrib.get("codelist")
                        if not existing.get("external_value_list") and attrib.get("external_value_list"):
                            existing["external_value_list"] = attrib.get("external_value_list")
                        if not existing.get("codelist_extensible") and attrib.get("codelist_extensible"):
                            existing["codelist_extensible"] = attrib.get("codelist_extensible")
                        _output[attrib.get("attribute_name")] = existing
                    else:
                        _output[attrib.get("attribute_name")] = attrib
                # connections = [
                #     connectors[cid]
                #     for cid in connectors
                #     if connectors[cid].start_object_id == _object_id
                # ]
                # for connection in connections:
                #     if connection.connector_type not in ("Association", "Generalization"):
                #         continue
                #     attrib = _output.setdefault(connection.name, {})
                #     if not attrib:
                #         attrib = dict(
                #             attribute_name=connection.name,
                #             attribute_type=objects[connection.end_object_id].name,
                #             attribute_cardinality=connection.source_card,
                #             note=None,
                #         )
                #     _output[connection.name] = attrib
                for _, attrib in enumerate(_output.values()):
                    _codelist = attrib.get("codelist")
                    if _codelist:
                        print("adding codelist: ", _codelist)
                        _cl = ct_content.get(_codelist)
                        if not _codelist.startswith("CNEW"):
                            _codelist_value = '=HYPERLINK("#\'{}\'!A2","{}")'.format(
                                _codelist, _codelist
                            )
                        else:
                            _codelist_value = "CNEW"
                    else:
                        _codelist_value = ""
                    write_cell(sheet, row=row_num, column=1, value=obj.name)
                    write_cell(
                        sheet, row=row_num, column=2, value=attrib["attribute_name"]
                    )
                    write_cell(
                        sheet, row=row_num, column=3, value=attrib["attribute_type"]
                    )
                    write_cell(
                        sheet,
                        row=row_num,
                        column=4,
                        value=attrib["attribute_cardinality"],
                    )
                    if attrib["attribute_note"]:
                        write_cell(
                            sheet, row=row_num, column=5, value=attrib["attribute_note"]
                        )
                    write_cell(
                        sheet, row=row_num, column=6, value=attrib.get("c_code", "")
                    )
                    write_cell(
                        sheet, row=row_num, column=7, value=attrib.get("pref_term", "")
                    )
                    write_cell(
                        sheet, row=row_num, column=8, value=attrib.get("definition", "")
                    )
                    write_cell(sheet, row=row_num, column=9, value=_codelist_value)
                    write_cell(
                        sheet,
                        row=row_num,
                        column=10,
                        value=attrib.get("external_value_list", ""),
                    )
                    write_cell(
                        sheet,
                        row=row_num,
                        column=11,
                        value=attrib.get("codelist_extensible", ""),
                    )
                    row_num += 1
    for c_code, _codelist in codelists.items():
        if c_code.strip().upper() == "CNEW":
            continue
        if _codelist is None:
            print("Missing codelist: ", c_code)
            continue
        _sheet = doc.create_sheet(c_code)
        for idx, colheader in enumerate(
            ["Code", "Preferred Term", "Synonyms", "Definition"], start=1
        ):
            write_cell(_sheet, row=1, column=idx, value=colheader, header=True)
        for idx, item in enumerate(_codelist.items, start=2):
            # type item: PermissibleValue
            write_cell(_sheet, row=idx, column=1, value=item.concept_c_code)
            write_cell(_sheet, row=idx, column=2, value=item.preferred_term)
            write_cell(
                _sheet,
                row=idx,
                column=3,
                value=";".join(item.synonyms) if item.synonyms else "",
            )
            write_cell(_sheet, row=idx, column=4, value=item.definition)
    # Span the column widths
    for wks in doc.worksheets:
        _columns = cols.get(wks.title, {})
        for colnum, length in _columns.items():
            column_width = min([length, 50])
            wks.column_dimensions[get_column_letter(colnum)].width = column_width

    # create the output directory if it doesn't exist
    if not os.path.exists(output_dir):
        os.mkdir(output_dir)
    for element in ("objects", "attributes", "connections"):
        if missing_references.get(element):
            for element_diff in missing_references.get(element):
                if isinstance(element_diff, tuple):
                    print(f"Missing {element}: {element_diff[0]} -> {element_diff[1]}")
                else:
                    print(f"Missing {element}: {element_diff}")
    fname = os.path.join(output_dir, f"{name}.xlsx")
    doc.save(fname)
    print(f"Generated USDM Excel file: {fname}")
